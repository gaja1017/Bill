from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from io import BytesIO


# Colors
LG_MAGENTA = "E6007E"
LIGHT_GRAY = "F5F5F5"
HEADER_BG = "F0F0F0"
TOTAL_BG = "FFE4E9"
BORDER_COLOR = "CCCCCC"
MEDIUM_GRAY = "808080"

thin_border = Border(
    left=Side(style='thin', color=BORDER_COLOR),
    right=Side(style='thin', color=BORDER_COLOR),
    top=Side(style='thin', color=BORDER_COLOR),
    bottom=Side(style='thin', color=BORDER_COLOR)
)

header_fill = PatternFill(start_color=HEADER_BG, end_color=HEADER_BG, fill_type="solid")
total_fill = PatternFill(start_color=TOTAL_BG, end_color=TOTAL_BG, fill_type="solid")
logo_fill = PatternFill(start_color=LIGHT_GRAY, end_color=LIGHT_GRAY, fill_type="solid")


def format_number(amount: int) -> str:
    if amount == 0:
        return "0"
    return f"{amount:,}"


def generate_quote_excel(quote_data: dict) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "견적서"

    # Column widths
    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 25
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 8
    ws.column_dimensions['F'].width = 14
    ws.column_dimensions['G'].width = 3
    ws.column_dimensions['H'].width = 12
    ws.column_dimensions['I'].width = 18

    current_row = 1

    # === HEADER SECTION ===
    # LG U+ Logo
    ws.merge_cells(f'A{current_row}:A{current_row+3}')
    ws[f'A{current_row}'] = "LG U+"
    ws[f'A{current_row}'].font = Font(size=18, bold=True, color=LG_MAGENTA)
    ws[f'A{current_row}'].alignment = Alignment(horizontal='center', vertical='center')
    ws[f'A{current_row}'].fill = logo_fill

    # QUOTATION
    ws[f'B{current_row}'] = "QUOTATION"
    ws[f'B{current_row}'].font = Font(size=20, color=MEDIUM_GRAY)

    # Manager info header
    ws[f'H{current_row}'] = "담당자명"
    ws[f'H{current_row}'].font = Font(size=8, bold=True)
    ws[f'H{current_row}'].border = thin_border
    ws[f'I{current_row}'] = quote_data.get('manager_name', '')
    ws[f'I{current_row}'].font = Font(size=8)
    ws[f'I{current_row}'].border = thin_border

    current_row += 1

    # Customer name
    customer_name = quote_data.get('customer_name', '')
    ws[f'B{current_row}'] = f"{customer_name} 귀하"
    ws[f'B{current_row}'].font = Font(size=12, bold=True)

    ws[f'H{current_row}'] = "연락처"
    ws[f'H{current_row}'].font = Font(size=8, bold=True)
    ws[f'H{current_row}'].border = thin_border
    ws[f'I{current_row}'] = quote_data.get('manager_phone', '')
    ws[f'I{current_row}'].font = Font(size=8)
    ws[f'I{current_row}'].border = thin_border

    current_row += 1

    ws[f'H{current_row}'] = "Fax"
    ws[f'H{current_row}'].font = Font(size=8, bold=True)
    ws[f'H{current_row}'].border = thin_border
    ws[f'I{current_row}'] = quote_data.get('manager_fax', '')
    ws[f'I{current_row}'].font = Font(size=8)
    ws[f'I{current_row}'].border = thin_border

    current_row += 1

    # Quote info
    ws[f'B{current_row}'] = "견적명"
    ws[f'B{current_row}'].font = Font(size=9)
    ws[f'C{current_row}'] = quote_data.get('quote_title', '')
    ws[f'C{current_row}'].font = Font(size=9)

    ws[f'H{current_row}'] = "E-mail"
    ws[f'H{current_row}'].font = Font(size=8, bold=True)
    ws[f'H{current_row}'].border = thin_border
    ws[f'I{current_row}'] = quote_data.get('manager_email', '')
    ws[f'I{current_row}'].font = Font(size=8)
    ws[f'I{current_row}'].border = thin_border

    current_row += 1

    ws[f'B{current_row}'] = "견적일"
    ws[f'B{current_row}'].font = Font(size=9)
    ws[f'C{current_row}'] = quote_data.get('quote_date', '')
    ws[f'C{current_row}'].font = Font(size=9)

    current_row += 1

    ws[f'B{current_row}'] = "참  조"
    ws[f'B{current_row}'].font = Font(size=9)
    ws[f'C{current_row}'] = customer_name
    ws[f'C{current_row}'].font = Font(size=9)

    current_row += 2

    # === MONTHLY PAYMENT SECTION ===
    # Section header
    ws.merge_cells(f'A{current_row}:D{current_row}')
    ws[f'A{current_row}'] = "1. 매월 납부 금액"
    ws[f'A{current_row}'].font = Font(size=10, bold=True)
    ws[f'A{current_row}'].fill = header_fill

    ws.merge_cells(f'E{current_row}:F{current_row}')
    ws[f'E{current_row}'] = "(단위 : 원/월, VAT 포함)"
    ws[f'E{current_row}'].font = Font(size=8, color=MEDIUM_GRAY)
    ws[f'E{current_row}'].alignment = Alignment(horizontal='right')
    ws[f'E{current_row}'].fill = header_fill

    current_row += 1

    # Table headers
    headers = ['품목', '서비스 종류', '서비스명', '단가', '수량', '금액']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=current_row, column=col, value=header)
        cell.font = Font(size=9, bold=True)
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal='center', vertical='center')

    current_row += 1

    monthly_total = 0
    products = quote_data.get("products", [])
    equipments = quote_data.get("equipments", [])

    # Filter out equipments with 0 price
    active_equipments = [eq for eq in equipments if eq.get("unit_price", 0) > 0]

    # Products - 서비스 항목 (VAT 포함 금액 그대로 사용)
    product_start_row = current_row
    num_products = len(products)
    for idx, product in enumerate(products):
        unit_price = product.get("unit_price", 0)  # VAT 포함 금액
        line_count = product.get("line_count", 1)
        speed = product.get("speed", "")
        ip_type = product.get("ip_type", "")
        amount = unit_price * line_count
        monthly_total += amount

        service_name = f"오피스넷 {ip_type} {speed}"

        if idx == 0:
            ws.cell(row=current_row, column=1, value="서비스").font = Font(size=9)
            ws.cell(row=current_row, column=2, value="오피스넷").font = Font(size=9)

        ws.cell(row=current_row, column=3, value=service_name).font = Font(size=9)
        ws.cell(row=current_row, column=4, value=format_number(unit_price)).font = Font(size=9)
        ws.cell(row=current_row, column=4).alignment = Alignment(horizontal='right')
        ws.cell(row=current_row, column=5, value=line_count).font = Font(size=9)
        ws.cell(row=current_row, column=5).alignment = Alignment(horizontal='center')
        ws.cell(row=current_row, column=6, value=format_number(amount)).font = Font(size=9)
        ws.cell(row=current_row, column=6).alignment = Alignment(horizontal='right')

        for col in range(1, 7):
            ws.cell(row=current_row, column=col).border = thin_border

        current_row += 1

    # Merge cells for products (서비스/오피스넷)
    if num_products > 1:
        ws.merge_cells(f'A{product_start_row}:A{product_start_row + num_products - 1}')
        ws.merge_cells(f'B{product_start_row}:B{product_start_row + num_products - 1}')

    # Equipment rentals - 통신장비 항목 (VAT 포함 금액 그대로 사용)
    equipment_start_row = current_row
    num_equipments = len(active_equipments)
    for idx, eq in enumerate(active_equipments):
        eq_name = eq.get("name", "")
        eq_price = eq.get("unit_price", 0)  # VAT 포함 금액
        eq_qty = eq.get("quantity", 1)
        eq_amount = eq_price * eq_qty
        monthly_total += eq_amount

        if idx == 0:
            ws.cell(row=current_row, column=1, value="통신장비").font = Font(size=9)
            ws.cell(row=current_row, column=2, value="장비임대").font = Font(size=9)

        ws.cell(row=current_row, column=3, value=eq_name).font = Font(size=9)
        ws.cell(row=current_row, column=4, value=format_number(eq_price)).font = Font(size=9)
        ws.cell(row=current_row, column=4).alignment = Alignment(horizontal='right')
        ws.cell(row=current_row, column=5, value=eq_qty).font = Font(size=9)
        ws.cell(row=current_row, column=5).alignment = Alignment(horizontal='center')
        ws.cell(row=current_row, column=6, value=format_number(eq_amount)).font = Font(size=9)
        ws.cell(row=current_row, column=6).alignment = Alignment(horizontal='right')

        for col in range(1, 7):
            ws.cell(row=current_row, column=col).border = thin_border

        current_row += 1

    # Merge cells for equipment (통신장비/장비임대)
    if num_equipments > 1:
        ws.merge_cells(f'A{equipment_start_row}:A{equipment_start_row + num_equipments - 1}')
        ws.merge_cells(f'B{equipment_start_row}:B{equipment_start_row + num_equipments - 1}')

    # 합계 row
    ws.cell(row=current_row, column=2, value="합계").font = Font(size=9, bold=True)
    ws.cell(row=current_row, column=2).alignment = Alignment(horizontal='center')
    ws.cell(row=current_row, column=6, value=format_number(monthly_total)).font = Font(size=9, bold=True, color=LG_MAGENTA)
    ws.cell(row=current_row, column=6).alignment = Alignment(horizontal='right')
    ws.cell(row=current_row, column=6).fill = total_fill

    for col in range(1, 7):
        ws.cell(row=current_row, column=col).border = thin_border

    current_row += 2

    # === ONE-TIME PAYMENT SECTION ===
    ws.merge_cells(f'A{current_row}:D{current_row}')
    ws[f'A{current_row}'] = "2. 일회성 납부 금액"
    ws[f'A{current_row}'].font = Font(size=10, bold=True)
    ws[f'A{current_row}'].fill = header_fill

    ws.merge_cells(f'E{current_row}:F{current_row}')
    ws[f'E{current_row}'] = "(단위 : 원, VAT 포함)"
    ws[f'E{current_row}'].font = Font(size=8, color=MEDIUM_GRAY)
    ws[f'E{current_row}'].alignment = Alignment(horizontal='right')
    ws[f'E{current_row}'].fill = header_fill

    current_row += 1

    # Table headers
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=current_row, column=col, value=header)
        cell.font = Font(size=9, bold=True)
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal='center', vertical='center')

    current_row += 1
    onetime_start_row = current_row

    one_time_charges = quote_data.get("one_time_charges", [])
    # Filter out charges with 0 price
    active_charges = [ch for ch in one_time_charges if ch.get("unit_price", 0) > 0]
    onetime_total = 0

    if active_charges:
        for idx, ch in enumerate(active_charges):
            ch_name = ch.get("name", "")
            display_name = "인터넷 설치비" if ch_name == "설치비" else ch_name
            service_type = "통신장비" if ch_name == "통신렉" else "설치비"

            ch_price = ch.get("unit_price", 0)  # VAT 포함 금액
            ch_qty = ch.get("quantity", 1)
            ch_amount = ch_price * ch_qty
            onetime_total += ch_amount

            if idx == 0:
                ws.cell(row=current_row, column=1, value="1회성 비용").font = Font(size=9)

            ws.cell(row=current_row, column=2, value=service_type).font = Font(size=9)
            ws.cell(row=current_row, column=3, value=display_name).font = Font(size=9)
            ws.cell(row=current_row, column=4, value=format_number(ch_price)).font = Font(size=9)
            ws.cell(row=current_row, column=4).alignment = Alignment(horizontal='right')
            ws.cell(row=current_row, column=5, value=ch_qty).font = Font(size=9)
            ws.cell(row=current_row, column=5).alignment = Alignment(horizontal='center')
            ws.cell(row=current_row, column=6, value=format_number(ch_amount)).font = Font(size=9)
            ws.cell(row=current_row, column=6).alignment = Alignment(horizontal='right')

            for col in range(1, 7):
                ws.cell(row=current_row, column=col).border = thin_border

            current_row += 1
    else:
        # No active charges - add a placeholder row
        ws.cell(row=current_row, column=1, value="1회성 비용").font = Font(size=9)
        ws.cell(row=current_row, column=2, value="-").font = Font(size=9)
        ws.cell(row=current_row, column=2).alignment = Alignment(horizontal='center')
        ws.cell(row=current_row, column=3, value="해당 없음").font = Font(size=9)
        ws.cell(row=current_row, column=4, value="0").font = Font(size=9)
        ws.cell(row=current_row, column=4).alignment = Alignment(horizontal='right')
        ws.cell(row=current_row, column=5, value="-").font = Font(size=9)
        ws.cell(row=current_row, column=5).alignment = Alignment(horizontal='center')
        ws.cell(row=current_row, column=6, value="0").font = Font(size=9)
        ws.cell(row=current_row, column=6).alignment = Alignment(horizontal='right')

        for col in range(1, 7):
            ws.cell(row=current_row, column=col).border = thin_border

        current_row += 1

    # Merge 품목 column
    if len(active_charges) > 1:
        ws.merge_cells(f'A{onetime_start_row}:A{onetime_start_row + len(active_charges) - 1}')

    # 합계 row
    ws.cell(row=current_row, column=2, value="합계").font = Font(size=9, bold=True)
    ws.cell(row=current_row, column=2).alignment = Alignment(horizontal='center')
    ws.cell(row=current_row, column=6, value=format_number(onetime_total)).font = Font(size=9, bold=True, color=LG_MAGENTA)
    ws.cell(row=current_row, column=6).alignment = Alignment(horizontal='right')
    ws.cell(row=current_row, column=6).fill = total_fill

    for col in range(1, 7):
        ws.cell(row=current_row, column=col).border = thin_border

    current_row += 2

    # === FOOTER ===
    ws.merge_cells(f'A{current_row}:F{current_row}')
    ws[f'A{current_row}'] = "오피스넷 : 중소기업, 프랜차이즈, 자영업자 등 SME 고객을 대상으로 고객 댁내까지 광케이블을 구성하여 100M부터 최대 10G의 속도를 제공하는 기업 인터넷 서비스 입니다."
    ws[f'A{current_row}'].font = Font(size=8)
    current_row += 1

    ws.merge_cells(f'A{current_row}:F{current_row}')
    ws[f'A{current_row}'] = "(유동IP는 단독사용 가능, 고정IP는 유동IP 1개 필수 사용)"
    ws[f'A{current_row}'].font = Font(size=8)
    current_row += 2

    # 특이사항
    ws[f'A{current_row}'] = "특 이 사 항"
    ws[f'A{current_row}'].font = Font(size=9, bold=True)
    current_row += 1

    notes = quote_data.get("notes", "")
    default_notes = [
        "※ 견적유효기간: 견적일로부터 1개월 이내",
        "※ 계약기간 : 3년 (신규설치비 면제, 광모뎀 임대료 면제, L2스위치 임대료 할인가 적용, 광모뎀/L2스위치 선택가능)",
        "※ 외곽지, 신규 건설현장 등은 초과 공사비가 별도 발생 하거나, 신규개통이 불가할 수 있습니다.",
        "※ 제반 사항은 LG유플러스 서비스 이용 약관에 따름",
        "※ 상기 금액은 VAT 포함 금액입니다."
    ]

    if notes:
        default_notes = notes.split('\n')

    for note in default_notes:
        if note.strip():
            ws.merge_cells(f'A{current_row}:F{current_row}')
            ws[f'A{current_row}'] = note.strip()
            ws[f'A{current_row}'].font = Font(size=8)
            current_row += 1

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
